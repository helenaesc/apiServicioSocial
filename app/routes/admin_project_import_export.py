from io import BytesIO
from flask import Blueprint, jsonify, request, send_file
from mysql.connector import Error
from openpyxl import load_workbook, Workbook
from ..database import execute_tx, fetch_all
from ..authz import require_role, ROLE_ADMIN

admin_project_import_export_bp = Blueprint("admin_project_import_export", __name__)

EXPECTED_HEADERS = [
    "general_name",
    "name",
    "partner_name",
    "modality_name",
    "week_days_name",
    "schedule_name",
    "slots",
    "schedule_description",
    "team_owners",
    "objectives",
    "activities",
    "clave",
    "competencies",
    "location",
    "duration",
    "audience",
    "max_hours",
    "comments",
]


def _clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value if value else None


def _int_or_none(value, field_name, min_value=0):
    if value is None or str(value).strip() == "":
        return None
    try:
        n = int(value)
    except:
        raise ValueError(f"{field_name} debe ser numérico")
    if n < min_value:
        raise ValueError(f"{field_name} debe ser >= {min_value}")
    return n


def _normalize_headers(row_values):
    return [_clean(v) for v in row_values]


def _validate_headers(headers):
    if headers != EXPECTED_HEADERS:
        return False
    return True


def _build_catalog_maps(cur):
    maps = {}

    cur.execute("SELECT id, name FROM partner")
    maps["partner"] = {str(r["name"]).strip().lower(): r["id"] for r in cur.fetchall()}

    cur.execute("SELECT id, name FROM modality")
    maps["modality"] = {str(r["name"]).strip().lower(): r["id"] for r in cur.fetchall()}

    cur.execute("SELECT id, name FROM week_days")
    maps["week_days"] = {str(r["name"]).strip().lower(): r["id"] for r in cur.fetchall()}

    cur.execute("SELECT id, name FROM schedule")
    maps["schedule"] = {str(r["name"]).strip().lower(): r["id"] for r in cur.fetchall()}

    return maps


def _lookup_catalog_id(value, map_dict, field_name):
    value_clean = _clean(value)
    if not value_clean:
        raise ValueError(f"{field_name} es obligatorio")
    found = map_dict.get(value_clean.lower())
    if not found:
        raise ValueError(f"{field_name} no existe: {value_clean}")
    return found


@admin_project_import_export_bp.post("/api/admin/projects/import")
def import_master_projects():
    role = require_role(request, {ROLE_ADMIN})
    if not role:
        return jsonify({"error": "No autorizado (solo ADMIN)"}), 401

    if "file" not in request.files:
        return jsonify({"error": "Debes subir un archivo Excel en el campo 'file'"}), 400

    uploaded_file = request.files["file"]
    if not uploaded_file or not uploaded_file.filename:
        return jsonify({"error": "Archivo inválido"}), 400

    filename = uploaded_file.filename.lower()
    if not filename.endswith(".xlsx"):
        return jsonify({"error": "Solo se permite formato .xlsx"}), 400

    try:
        workbook = load_workbook(uploaded_file, data_only=True)
    except Exception as e:
        return jsonify({"error": f"No se pudo leer el archivo Excel: {str(e)}"}), 400

    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return jsonify({"error": "El archivo está vacío"}), 400

    headers = _normalize_headers(rows[0])
    if not _validate_headers(headers):
        return jsonify({
            "error": "Encabezados inválidos",
            "expected_headers": EXPECTED_HEADERS,
            "received_headers": headers
        }), 400

    data_rows = rows[1:]
    if not data_rows:
        return jsonify({"error": "No hay filas para importar"}), 400

    def tx(conn, cur):
        catalog_maps = _build_catalog_maps(cur)

        inserted = 0
        failed = 0
        errors = []

        for excel_row_number, row in enumerate(data_rows, start=2):
            try:
                row_dict = dict(zip(EXPECTED_HEADERS, row))

                general_name = _clean(row_dict["general_name"])
                name = _clean(row_dict["name"])
                partner_name = _clean(row_dict["partner_name"])
                modality_name = _clean(row_dict["modality_name"])
                week_days_name = _clean(row_dict["week_days_name"])
                schedule_name = _clean(row_dict["schedule_name"])

                if not general_name:
                    raise ValueError("general_name es obligatorio")
                if not name:
                    raise ValueError("name es obligatorio")
                if not partner_name:
                    raise ValueError("partner_name es obligatorio")
                if not modality_name:
                    raise ValueError("modality_name es obligatorio")
                if not week_days_name:
                    raise ValueError("week_days_name es obligatorio")
                if not schedule_name:
                    raise ValueError("schedule_name es obligatorio")

                id_partner = _lookup_catalog_id(partner_name, catalog_maps["partner"], "partner_name")
                id_modality = _lookup_catalog_id(modality_name, catalog_maps["modality"], "modality_name")
                id_week_days = _lookup_catalog_id(week_days_name, catalog_maps["week_days"], "week_days_name")
                id_schedule = _lookup_catalog_id(schedule_name, catalog_maps["schedule"], "schedule_name")

                slots = _int_or_none(row_dict["slots"], "slots", 0)
                if slots is None:
                    slots = 0

                max_hours = _int_or_none(row_dict["max_hours"], "max_hours", 0)

                schedule_description = _clean(row_dict["schedule_description"])
                team_owners = _clean(row_dict["team_owners"])
                objectives = _clean(row_dict["objectives"])
                activities = _clean(row_dict["activities"])
                clave = _clean(row_dict["clave"])
                competencies = _clean(row_dict["competencies"])
                location = _clean(row_dict["location"])
                duration = _clean(row_dict["duration"])
                audience = _clean(row_dict["audience"])
                comments = _clean(row_dict["comments"])

                # regla de duplicado
                cur.execute(
                    """
                    SELECT id
                    FROM project
                    WHERE COALESCE(general_name, '') = COALESCE(%s, '')
                      AND COALESCE(name, '') = COALESCE(%s, '')
                      AND COALESCE(clave, '') = COALESCE(%s, '')
                    LIMIT 1
                    """,
                    [general_name, name, clave]
                )
                if cur.fetchone():
                    raise ValueError("Proyecto duplicado por general_name + name + clave")

                cur.execute(
                    """
                    INSERT INTO project
                    (
                        general_name,
                        name,
                        id_partner,
                        id_modality,
                        id_week_days,
                        id_schedule,
                        slots,
                        schedule_description,
                        team_owners,
                        objectives,
                        activities,
                        clave,
                        competencies,
                        location,
                        duration,
                        audience,
                        max_hours,
                        comments
                    )
                    VALUES
                    (
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
                    )
                    """,
                    [
                        general_name,
                        name,
                        id_partner,
                        id_modality,
                        id_week_days,
                        id_schedule,
                        slots,
                        schedule_description,
                        team_owners,
                        objectives,
                        activities,
                        clave,
                        competencies,
                        location,
                        duration,
                        audience,
                        max_hours,
                        comments,
                    ]
                )

                inserted += 1

            except Exception as e:
                failed += 1
                errors.append({
                    "row": excel_row_number,
                    "error": str(e)
                })

        return {
            "status": 200,
            "message": "Importación completada",
            "inserted": inserted,
            "failed": failed,
            "errors": errors
        }

    try:
        result = execute_tx(tx)
        return jsonify(result), 200
    except Error as e:
        return jsonify({"error": f"Error de base de datos: {e.msg}"}), 500
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@admin_project_import_export_bp.get("/api/admin/projects/export")
def export_master_projects():
    role = require_role(request, {ROLE_ADMIN})
    if not role:
        return jsonify({"error": "No autorizado (solo ADMIN)"}), 401

    try:
        rows = fetch_all(
            """
            SELECT
                p.general_name,
                p.name,
                pa.name AS partner_name,
                m.name AS modality_name,
                wd.name AS week_days_name,
                sc.name AS schedule_name,
                p.slots,
                p.schedule_description,
                p.team_owners,
                p.objectives,
                p.activities,
                p.clave,
                p.competencies,
                p.location,
                p.duration,
                p.audience,
                p.max_hours,
                p.comments
            FROM project p
            LEFT JOIN partner pa ON pa.id = p.id_partner
            LEFT JOIN modality m ON m.id = p.id_modality
            LEFT JOIN week_days wd ON wd.id = p.id_week_days
            LEFT JOIN schedule sc ON sc.id = p.id_schedule
            ORDER BY p.general_name, p.name, p.id
            """
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "projects"

        ws.append(EXPECTED_HEADERS)

        for row in rows:
            ws.append([
                row.get("general_name"),
                row.get("name"),
                row.get("partner_name"),
                row.get("modality_name"),
                row.get("week_days_name"),
                row.get("schedule_name"),
                row.get("slots"),
                row.get("schedule_description"),
                row.get("team_owners"),
                row.get("objectives"),
                row.get("activities"),
                row.get("clave"),
                row.get("competencies"),
                row.get("location"),
                row.get("duration"),
                row.get("audience"),
                row.get("max_hours"),
                row.get("comments"),
            ])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name="projects_export.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Error as e:
        return jsonify({"error": f"Error de base de datos: {e.msg}"}), 500
    except Exception as e:
        return jsonify({"error": str(e)}), 500